from pathlib import Path
from langchain_core.documents import Document

supported_extensions = [".txt", ".pdf", ".docx", ".csv", ".json", ".md"]

def load_document(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in supported_extensions:
        raise ValueError(f"Unsupported file type: {extension}")

    if extension in [".txt", ".md"]:
        content = path.read_text(encoding="utf-8")

    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        content = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif extension == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        content = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    elif extension == ".csv":
        import pandas as pd

        content = pd.read_csv(path).to_string(index=False)

    elif extension == ".json":
        import json

        content = json.dumps(
            json.loads(path.read_text(encoding="utf-8")),
            indent=2
        )

    if not content.strip():
        raise ValueError(f"No readable text found in {path.name}")

    return Document(
        page_content=content,
        metadata={
            "source": path.name,
            "file_path": str(path.resolve()),
            "file_type": extension
        }
    )

def load_documents(directory):
    directory = Path(directory)
    documents = []

    for path in directory.rglob("*"):
        if path.is_file() and path.suffix.lower() in supported_extensions:
            try:
                document = load_document(path)
                documents.append(document)
            except Exception as e:
                print(f"Skipping {path.name}: {e}")

    return documents